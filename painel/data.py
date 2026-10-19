"""Acesso à persistência local e importação da aba GERENCIAL do Excel."""
import json
import secrets
from contextlib import contextmanager, closing
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict

from .config import AREAS, Area, EXCEL_ROW_MAP, destination_channel
from .metrics import number


@dataclass
class Snapshot:
    period: str
    reference_date: str
    source: str
    areas: Dict[str, dict] = field(default_factory=dict)
    errors: Dict[str, dict] = field(default_factory=dict)
    file_path: str = ''
    imported_at: str = ''

    def to_dict(self):
        return {
            'period': self.period,
            'reference_date': self.reference_date,
            'source': self.source,
            'areas': self.areas,
            'errors': self.errors,
            'file_path': self.file_path,
            'imported_at': self.imported_at,
        }

    @classmethod
    def from_dict(cls, data):
        return cls(**{k: data.get(k, default) for k, default in {
            'period': '', 'reference_date': '', 'source': '', 'areas': {},
            'errors': {}, 'file_path': '', 'imported_at': '',
        }.items()})


def app_data_dir():
    """Não grava dados de operação dentro de Program Files ou do ZIP."""
    if sys.platform == 'win32':
        import os
        base = Path(os.environ.get('LOCALAPPDATA', Path.home())) / 'PainelGestaoIntegrada'
    else:
        base = Path.home() / '.painel_gestao_integrada'
    base.mkdir(parents=True, exist_ok=True)
    return base


class SQLiteSnapshotRepository:
    def __init__(self, path=None):
        self.path = Path(path or app_data_dir() / 'painel.db')
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self._conn() as conn:
            conn.execute('''CREATE TABLE IF NOT EXISTS snapshots (
                period TEXT PRIMARY KEY,
                imported_at TEXT NOT NULL,
                source TEXT NOT NULL,
                payload TEXT NOT NULL
            )''')
            conn.execute('''CREATE TABLE IF NOT EXISTS daily_entries (
                reference_date TEXT NOT NULL,
                area_key TEXT NOT NULL,
                payload TEXT NOT NULL,
                saved_at TEXT NOT NULL,
                PRIMARY KEY (reference_date, area_key)
            )''')
            conn.execute('''CREATE TABLE IF NOT EXISTS settings (
                setting_key TEXT PRIMARY KEY,
                setting_value TEXT NOT NULL
            )''')
            # V5: catálogo estável de setores. Ocultar não altera o histórico.
            conn.execute('''CREATE TABLE IF NOT EXISTS area_catalog (
                area_key TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                short TEXT NOT NULL COLLATE NOCASE UNIQUE,
                primary_label TEXT NOT NULL,
                secondary_label TEXT NOT NULL,
                excel_col TEXT NOT NULL DEFAULT '',
                is_visible INTEGER NOT NULL DEFAULT 1 CHECK(is_visible IN (0,1)),
                sort_index INTEGER NOT NULL
            )''')
            for order, area in enumerate(AREAS):
                conn.execute('''INSERT OR IGNORE INTO area_catalog
                    (area_key,title,short,primary_label,secondary_label,excel_col,is_visible,sort_index)
                    VALUES (?,?,?,?,?,?,?,?)''',
                    (area.key, area.title, area.short, area.primary, area.secondary,
                     area.excel_col, 0 if area.key == 'garantia' else 1, order))

    @contextmanager
    def _conn(self):
        """Conexões são fechadas após cada transação, inclusive Python 3.13+."""
        conn = sqlite3.connect(self.path, timeout=30)
        try:
            with conn:
                yield conn
        finally:
            conn.close()

    def put(self, snap: Snapshot):
        snap.imported_at = datetime.now().isoformat(timespec='seconds')
        with self._conn() as conn:
            conn.execute('''INSERT INTO snapshots(period, imported_at, source, payload)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(period) DO UPDATE SET
                    imported_at=excluded.imported_at,
                    source=excluded.source,
                    payload=excluded.payload''',
                (snap.period, snap.imported_at, snap.source, json.dumps(snap.to_dict(), ensure_ascii=False)))

    def get(self, period):
        with self._conn() as conn:
            row = conn.execute('SELECT payload FROM snapshots WHERE period=?', (period,)).fetchone()
        return Snapshot.from_dict(json.loads(row[0])) if row else None

    def periods(self):
        with self._conn() as conn:
            rows = conn.execute('SELECT period FROM snapshots ORDER BY period DESC').fetchall()
        return [r[0] for r in rows]

    @staticmethod
    def _to_area(row):
        return Area(row[0], row[1], row[2], row[3], row[4], row[5])

    def list_areas(self, include_hidden=False):
        """Catálogo ordenado; IDs não mudam após renomeações."""
        with self._conn() as conn:
            rows = conn.execute('SELECT area_key,title,short,primary_label,secondary_label,excel_col '
                                'FROM area_catalog' + ('' if include_hidden else ' WHERE is_visible=1') +
                                ' ORDER BY sort_index,area_key').fetchall()
        return [self._to_area(r) for r in rows]

    def list_area_entries(self):
        with self._conn() as conn:
            rows = conn.execute('SELECT area_key,title,short,primary_label,secondary_label,excel_col,is_visible '
                                'FROM area_catalog ORDER BY sort_index,area_key').fetchall()
        return [{'area': self._to_area(r[:6]), 'visible': bool(r[6])} for r in rows]

    def get_area(self, area_key):
        with self._conn() as conn:
            row = conn.execute('SELECT area_key,title,short,primary_label,secondary_label,excel_col '
                               'FROM area_catalog WHERE area_key=?', (area_key,)).fetchone()
        return self._to_area(row) if row else None

    def is_area_visible(self, area_key):
        with self._conn() as conn:
            row = conn.execute('SELECT is_visible FROM area_catalog WHERE area_key=?',
                               (area_key,)).fetchone()
        return bool(row[0]) if row else False

    @staticmethod
    def _validate_area(title, short, primary, secondary):
        title, short, primary, secondary = [str(s).strip() for s in (title, short, primary, secondary)]
        if not title or not short or not primary:
            raise ValueError('Preencha nome completo, nome curto e indicador principal.')
        if len(title) > 100 or len(short) > 55 or len(primary) > 100 or len(secondary) > 100:
            raise ValueError('Nome excede o limite permitido (100 caracteres; curto: 55).')
        return title, short, primary, secondary

    def create_area(self, title, short, primary='Previsto / Produção (Qtd)',
                    secondary='Realizado (Qtd)'):
        """Novo painel começa visível, vazio, com a mesma estrutura de métricas."""
        title, short, primary, secondary = self._validate_area(title, short, primary, secondary)
        key = 'custom_' + secrets.token_hex(8)
        try:
            with self._conn() as conn:
                pos = conn.execute('SELECT COALESCE(MAX(sort_index),-1)+1 FROM area_catalog').fetchone()[0]
                conn.execute('INSERT INTO area_catalog '
                             '(area_key,title,short,primary_label,secondary_label,excel_col,is_visible,sort_index) '
                             'VALUES (?,?,?,?,?,?,1,?)', (key,title,short,primary,secondary,'',pos))
        except sqlite3.IntegrityError as exc:
            raise ValueError('Já existe um painel com esse nome curto.') from exc
        return self.get_area(key)

    def rename_area(self, key, title, short, primary, secondary):
        title, short, primary, secondary = self._validate_area(title, short, primary, secondary)
        try:
            with self._conn() as conn:
                cursor = conn.execute('UPDATE area_catalog SET title=?, short=?, primary_label=?, secondary_label=? '
                                      'WHERE area_key=?', (title,short,primary,secondary,key))
                if cursor.rowcount != 1:
                    raise ValueError('Painel não encontrado.')
        except sqlite3.IntegrityError as exc:
            raise ValueError('Já existe um painel com esse nome curto.') from exc
        return self.get_area(key)

    def set_area_visible(self, key, visible):
        with self._conn() as conn:
            cursor = conn.execute('UPDATE area_catalog SET is_visible=? WHERE area_key=?',
                                  (int(bool(visible)),key))
            if cursor.rowcount != 1:
                raise ValueError('Painel não encontrado.')
        # Nenhum dado mensal, diário ou métrica é excluído.

    def get_setting(self, key, default=None):
        with self._conn() as conn:
            row = conn.execute('SELECT setting_value FROM settings WHERE setting_key=?', (key,)).fetchone()
        return json.loads(row[0]) if row else default

    def set_setting(self, key, value):
        with self._conn() as conn:
            conn.execute('INSERT INTO settings(setting_key, setting_value) VALUES (?, ?) '
                         'ON CONFLICT(setting_key) DO UPDATE SET setting_value=excluded.setting_value',
                         (key, json.dumps(value, ensure_ascii=False)))

    def available_days(self, period):
        with self._conn() as conn:
            rows = conn.execute('SELECT DISTINCT reference_date FROM daily_entries '
                                'WHERE reference_date LIKE ? ORDER BY reference_date DESC',
                                (period + '-%',)).fetchall()
        return [r[0] for r in rows]

    def get_day(self, reference_date):
        """Somente registros efetivamente lançados naquele dia (sem usar dados futuros)."""
        ref = date.fromisoformat(reference_date)
        with self._conn() as conn:
            rows = conn.execute('SELECT area_key, payload FROM daily_entries WHERE reference_date=?',
                                (ref.isoformat(),)).fetchall()
        if not rows:
            return None
        return Snapshot(period=ref.strftime('%Y-%m'), reference_date=reference_date,
                        source='LANÇAMENTOS DO DIA',
                        areas={key: json.loads(payload) for key, payload in rows})

    def save_area_with_day(self, monthly: Snapshot, day, area_key, values):
        """Salva histórico diário e consolidação do último dia da área de modo atômico."""
        now = datetime.now().isoformat(timespec='seconds')
        monthly.imported_at = now
        with self._conn() as conn:
            latest = conn.execute('SELECT max(reference_date) FROM daily_entries '
                                  'WHERE area_key=? AND reference_date LIKE ?',
                                  (area_key, monthly.period + '-%')).fetchone()[0]
            conn.execute('INSERT INTO daily_entries(reference_date,area_key,payload,saved_at) '
                         'VALUES (?,?,?,?) ON CONFLICT(reference_date,area_key) '
                         'DO UPDATE SET payload=excluded.payload,saved_at=excluded.saved_at',
                         (day, area_key, json.dumps(values, ensure_ascii=False), now))
            # Um lançamento antigo não deve substituir a visão consolidada mais recente.
            current = conn.execute('SELECT payload FROM snapshots WHERE period=?',
                                   (monthly.period,)).fetchone()
            old = Snapshot.from_dict(json.loads(current[0])) if current else None
            if latest is None or day >= latest or old is None:
                conn.execute('INSERT INTO snapshots(period,imported_at,source,payload) '
                             'VALUES (?,?,?,?) ON CONFLICT(period) DO UPDATE SET '
                             'imported_at=excluded.imported_at,source=excluded.source,payload=excluded.payload',
                             (monthly.period, now, monthly.source,
                              json.dumps(monthly.to_dict(), ensure_ascii=False)))

    def _backup(self):
        """Backup íntegro do SQLite antes de operações destrutivas."""
        backup_dir = self.path.parent / 'backups'
        backup_dir.mkdir(parents=True, exist_ok=True)
        name = 'PGI_antes_limpeza_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.db'
        target = backup_dir / name
        with self._conn() as source, closing(sqlite3.connect(target)) as dest:
            source.backup(dest)
        return target

    def clear_area_period(self, period, area_key):
        if self.get_area(area_key) is None:
            raise ValueError('Área inválida')
        backup = self._backup()
        with self._conn() as conn:
            row = conn.execute('SELECT payload FROM snapshots WHERE period=?', (period,)).fetchone()
            if row:
                snap = Snapshot.from_dict(json.loads(row[0]))
                snap.areas.pop(area_key, None)
                snap.errors.pop(area_key, None)
                snap.source = 'EDIÇÃO LOCAL • LIMPEZA'
                snap.file_path = ''
                snap.imported_at = datetime.now().isoformat(timespec='seconds')
                conn.execute('UPDATE snapshots SET imported_at=?,source=?,payload=? WHERE period=?',
                             (snap.imported_at, snap.source, json.dumps(snap.to_dict(), ensure_ascii=False), period))
            conn.execute('DELETE FROM daily_entries WHERE area_key=? AND reference_date LIKE ?',
                         (area_key, period + '-%'))
        return backup

    def clear_day_area(self, reference_date, area_key):
        if self.get_area(area_key) is None:
            raise ValueError('Área inválida')
        period = date.fromisoformat(reference_date).strftime('%Y-%m')
        backup = self._backup()
        with self._conn() as conn:
            conn.execute('DELETE FROM daily_entries WHERE reference_date=? AND area_key=?',
                         (reference_date, area_key))
            row = conn.execute('SELECT payload FROM snapshots WHERE period=?', (period,)).fetchone()
            if row:
                snap = Snapshot.from_dict(json.loads(row[0]))
                latest = conn.execute('SELECT reference_date,payload FROM daily_entries '
                                      'WHERE area_key=? AND reference_date LIKE ? '
                                      'ORDER BY reference_date DESC LIMIT 1',
                                      (area_key, period + '-%')).fetchone()
                if latest:
                    snap.areas[area_key] = json.loads(latest[1])
                    snap.reference_date = latest[0]
                else:
                    snap.areas.pop(area_key, None)
                    snap.errors.pop(area_key, None)
                snap.source = 'EDIÇÃO LOCAL • LIMPEZA'
                snap.file_path = ''
                snap.imported_at = datetime.now().isoformat(timespec='seconds')
                conn.execute('UPDATE snapshots SET imported_at=?,source=?,payload=? WHERE period=?',
                             (snap.imported_at, snap.source, json.dumps(snap.to_dict(), ensure_ascii=False), period))
        return backup

    def replace_month_from_excel(self, snap):
        """Reimporta as áreas fixas preservando os painéis personalizados.

        A planilha original não tem colunas para os painéis criados manualmente.
        Seus históricos diários e consolidados não podem ser substituídos.
        """
        if self.get(snap.period) or self.available_days(snap.period):
            self._backup()
        snap.imported_at = datetime.now().isoformat(timespec='seconds')
        with self._conn() as conn:
            fixed_keys = {a.key for a in AREAS}
            previous = conn.execute('SELECT payload FROM snapshots WHERE period=?',
                                    (snap.period,)).fetchone()
            if previous:
                current = Snapshot.from_dict(json.loads(previous[0]))
                for key, metrics in current.areas.items():
                    if key not in fixed_keys:
                        snap.areas[key] = metrics
                for key, errors in current.errors.items():
                    if key not in fixed_keys:
                        snap.errors[key] = errors
            # Reimportar somente o modelo Excel, sem tocar nos registros customizados.
            fixed = tuple(sorted(fixed_keys))
            placeholders = ','.join('?' for _ in fixed)
            conn.execute('DELETE FROM daily_entries WHERE reference_date LIKE ? '
                         f'AND area_key IN ({placeholders})', (snap.period + '-%', *fixed))
            conn.execute('INSERT INTO snapshots(period,imported_at,source,payload) '
                         'VALUES (?,?,?,?) ON CONFLICT(period) DO UPDATE SET '
                         'imported_at=excluded.imported_at,source=excluded.source,payload=excluded.payload',
                         (snap.period, snap.imported_at, snap.source,
                          json.dumps(snap.to_dict(), ensure_ascii=False)))

    def clear_all_data(self):
        """Exclui lançamentos e competências; mantém preferências e faz backup."""
        backup = self._backup()
        with self._conn() as conn:
            conn.execute('DELETE FROM daily_entries')
            conn.execute('DELETE FROM snapshots')
            conn.execute('INSERT INTO settings(setting_key,setting_value) VALUES (?,?) '
                         'ON CONFLICT(setting_key) DO UPDATE SET setting_value=excluded.setting_value',
                         ('first_run_initialized', 'true'))
        return backup

    def seed_if_empty(self, example_path):
        # Evita ressuscitar dados de demonstração após a limpeza geral.
        if self.get_setting('first_run_initialized', False):
            return
        if not self.periods() and Path(example_path).is_file():
            example = json.loads(Path(example_path).read_text(encoding='utf-8'))
            self.put(Snapshot.from_dict(example))
        self.set_setting('first_run_initialized', True)


class ExcelGerencialAdapter:
    """Mapeamento EXPLÍCITO da aba GERENCIAL fornecida pelo usuário.

    Excel só fornece um valor por área/linha. Exibimos nas subcolunas visuais
    conforme o indicador principal ou consolidado, sem duplicar o número.
    Fórmulas com erro/sem cache não viram zero.
    """

    def import_file(self, file_path: str) -> Snapshot:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('Instale a dependência openpyxl: pip install -r requirements.txt') from exc

        path = Path(file_path).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(path)

        workbook = load_workbook(str(path), data_only=True, read_only=True)
        try:
            if 'GERENCIAL' not in workbook.sheetnames:
                raise ValueError('A planilha precisa ter a aba GERENCIAL.')
            sheet = workbook['GERENCIAL']
            raw_date = sheet['C1'].value
            if isinstance(raw_date, datetime):
                reference = raw_date.date()
            elif isinstance(raw_date, date):
                reference = raw_date
            elif isinstance(raw_date, str):
                try:
                    reference = date.fromisoformat(raw_date[:10])
                except ValueError as exc:
                    raise ValueError('Data de referência inválida em GERENCIAL!C1.') from exc
            else:
                raise ValueError('Data de referência ausente em GERENCIAL!C1; revise a origem.')

            snapshot = Snapshot(
                period=reference.strftime('%Y-%m'),
                reference_date=reference.isoformat(),
                source=f'EXCEL: {path.name}',
                file_path=str(path),
            )
            for area in AREAS:
                metrics = {}
                errors = {}
                for key, row_index in EXCEL_ROW_MAP.items():
                    cell = sheet[f'{area.excel_col}{row_index}']
                    raw = cell.value
                    ch = destination_channel(area, key)
                    metrics[key] = {ch: number(raw)}
                    if isinstance(raw, str) and raw.startswith('#'):
                        errors[key] = f'{cell.coordinate}: {raw}'
                ch = destination_channel(area, 'percentual')
                attainment = metrics['percentual'].get(ch)
                metrics['falta'] = {ch: max(0., 1 - attainment) if attainment is not None else None}
                snapshot.areas[area.key] = metrics
                if errors:
                    snapshot.errors[area.key] = errors
            return snapshot
        finally:
            workbook.close()


class SAPDataProvider:
    """Contrato de integração futura (sem fingir que SAP está conectado)."""

    def fetch_snapshot(self, reference: date) -> Snapshot:
        raise NotImplementedError(
            'Implementar autenticação, consultas autorizadas e mapeamento SAP nesta classe.'
        )
